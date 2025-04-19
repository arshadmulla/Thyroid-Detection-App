import streamlit as st
import pickle
import os
from docx import Document
from datetime import datetime
import pandas as pd
import matplotlib.pyplot as plt
import smtplib
from email.message import EmailMessage
from fpdf import FPDF
from PIL import Image

# Load model
model_path = 'model.pkl'
if not os.path.exists(model_path):
    st.error("Model file not found.")
else:
    with open(model_path, 'rb') as f:
        model = pickle.load(f)

diagnoses = {0: 'Negative', 1: 'Hypothyroid', 2: 'Hyperthyroid'}
diagnosis_color = '#F63366'
title_color = '#F63366'

def preprocess_inputs(*args):
    binary_map = {'Yes': 1, 'No': 0, '': None}
    age, sex, on_thyroxine, query_on_thyroxine, on_antithyroid_meds, sick, pregnant, \
    thyroid_surgery, I131_treatment, query_hypothyroid, query_hyperthyroid, lithium, \
    goitre, tumor, hypopituitary, psych, TSH, T3, TT4, T4U, FTI = args

    sex = 1 if sex == 'F' else 0 if sex == 'M' else None
    flags = [on_thyroxine, query_on_thyroxine, on_antithyroid_meds, sick, pregnant,
             thyroid_surgery, I131_treatment, query_hypothyroid, query_hyperthyroid,
             lithium, goitre, tumor, hypopituitary, psych]
    flags = [binary_map.get(flag, None) for flag in flags]

    return [age, sex] + flags + [TSH, T3, TT4, T4U, FTI]

def predict_diagnosis_with_confidence(inputs):
    prediction = model.predict([inputs])[0]
    confidence = max(model.predict_proba([inputs])[0]) * 100
    return prediction, confidence

def generate_pdf(name, diagnosis, confidence, conclusion):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, 'Thyroid Diagnosis Report', ln=True, align='C')
    pdf.ln(10)
    pdf.set_font("Arial", '', 12)
    pdf.cell(0, 10, f"Patient Name: {name}", ln=True)
    pdf.cell(0, 10, f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=True)
    pdf.cell(0, 10, f"Diagnosis: {diagnosis}", ln=True)
    pdf.cell(0, 10, f"Confidence: {confidence:.2f}%", ln=True)
    pdf.multi_cell(0, 10, f"Conclusion: {conclusion}")

    # Add diagnosis image
    image_path = ""
    if diagnosis == "Negative":
        image_path = "images/healthy.png"
    elif diagnosis == "Hypothyroid":
        image_path = "images/hypothyroid.png"
    elif diagnosis == "Hyperthyroid":
        image_path = "images/hyperthyroid.png"

    if image_path and os.path.exists(image_path):
        pdf.image(image_path, x=10, y=pdf.get_y() + 10, w=180)

    filename = f"{name.replace(' ', '_')}_thyroid_report.pdf"
    pdf.output(filename)
    return filename

def send_email_report(to_email, report_file):
    try:
        sender = "salmamulla12@gmail.com"
        password = "pgmf pmka yqdm twlh"
        msg = EmailMessage()
        msg['Subject'] = 'Thyroid Diagnosis Report'
        msg['From'] = sender
        msg['To'] = to_email
        msg.set_content('Attached is the thyroid diagnosis report.')

        with open(report_file, 'rb') as f:
            file_data = f.read()
            file_name = os.path.basename(report_file)
            msg.add_attachment(file_data, maintype='application', subtype='pdf', filename=file_name)

        with smtplib.SMTP_SSL('smtp.gmail.com', 465) as smtp:
            smtp.login(sender, password)
            smtp.send_message(msg)
        return True
    except Exception as e:
        st.error(f"Email failed to send: {e}")
        return False

def main():
    st.markdown("""
    <style>
    .main {
        background-image: url('https://images.unsplash.com/photo-1588776814546-bfc763b55c2e');
        background-size: cover;
        padding: 20px;
        border-radius: 10px;
    }
    h1 {
        color: white;
        text-shadow: 2px 2px 5px #000;
    }
    .stButton button {
        background-color: #F63366;
        color: white;
        border-radius: 10px;
        padding: 10px 24px;
        transition: 0.4s;
    }
    .stButton button:hover {
        background-color: #a61d3b;
        transform: scale(1.05);
    }
    </style>
    <div class='main'>
        <h1>Thyroid Diagnosis Predictor</h1>
        <p style='color:white;'>An intelligent tool to predict thyroid-related diseases using AI.</p>
    </div>
""", unsafe_allow_html=True)

    st.sidebar.markdown("<h1 style='color: #F63366;'>Salma Mulla</h1>", unsafe_allow_html=True)
    st.sidebar.markdown("""<p>This is a smart AI-based system to detect thyroid conditions using patient data inputs. 
        It provides instant results and sends a personalized report via email.</p>
        <p><strong>About Project:</strong><br>
        Developed using Streamlit, scikit-learn, and email automation, this project predicts thyroid disorders using clinical inputs and delivers customized reports.</p>
    """, unsafe_allow_html=True)

    with st.expander("ℹ️ Input Attribute Descriptions"):
        st.markdown("""
- **Age**: Age of the patient  
- **Sex**: Gender (M/F)  
- **On Thyroxine**: Currently on thyroxine medication (Yes/No)  
- **Query on Thyroxine**: Uncertain about being on thyroxine  
- **On Antithyroid Meds**: Taking anti-thyroid medication  
- **Sick**: Patient currently sick  
- **Pregnant**: Patient is pregnant  
- **Thyroid Surgery**: History of thyroid surgery  
- **I131 Treatment**: Radioiodine treatment history  
- **Query Hypothyroid**: Suspected hypothyroidism  
- **Query Hyperthyroid**: Suspected hyperthyroidism  
- **Lithium**: Taking lithium medication  
- **Goitre**: Enlarged thyroid gland  
- **Tumor**: Thyroid-related tumor  
- **Hypopituitary**: Pituitary gland dysfunction  
- **Psych**: Psychological disorders  
- **TSH**: Thyroid-stimulating hormone level  
- **T3**: Triiodothyronine hormone level  
- **TT4**: Total thyroxine level  
- **T4U**: Thyroxine uptake  
- **FTI**: Free thyroxine index  
""")

    st.markdown("### Enter Patient Information")
    name = st.text_input("Patient Name")
    email = st.text_input("Recipient Email")

    col1, col2, col3 = st.columns(3)
    with col1:
        age = st.number_input('Age', value=None)
        query_on_thyroxine = st.selectbox('Query On Thyroxine', ['','No','Yes'])
        pregnant = st.selectbox('Pregnant', ['','No','Yes'])
        query_hypothyroid = st.selectbox('Query Hypothyroid', ['','No','Yes'])
        goitre = st.selectbox('Goitre', ['','No','Yes'])
        psych = st.selectbox('Psych', ['','No','Yes'])
        TT4 = st.number_input('TT4', value=None)
    with col2:
        sex = st.selectbox('Sex', ['','M','F'])
        on_antithyroid_meds = st.selectbox('On Antithyroid Meds', ['','No','Yes'])
        thyroid_surgery = st.selectbox('Thyroid Surgery', ['','No','Yes'])
        query_hyperthyroid = st.selectbox('Query Hyperthyroid', ['','No','Yes'])
        tumor = st.selectbox('Tumor', ['','No','Yes'])
        TSH = st.number_input('TSH', value=None)
        T4U = st.number_input('T4U', value=None)
    with col3:
        on_thyroxine = st.selectbox('On Thyroxine', ['','No','Yes'])
        sick = st.selectbox('Sick', ['','No','Yes'])
        I131_treatment = st.selectbox('I131 Treatment', ['','No','Yes'])
        lithium = st.selectbox('Lithium', ['','No','Yes'])
        hypopituitary = st.selectbox('Hypopituitary', ['','No','Yes'])
        T3 = st.number_input('T3', value=None)
        FTI = st.number_input('FTI', value=None)

    if st.button('Detect'):
        if not name or not email:
            st.error("Please enter both patient name and recipient email.")
        else:
            inputs = preprocess_inputs(age, sex, on_thyroxine, query_on_thyroxine, on_antithyroid_meds, sick, pregnant, thyroid_surgery, I131_treatment, query_hypothyroid, query_hyperthyroid, lithium, goitre, tumor, hypopituitary, psych, TSH, T3, TT4, T4U, FTI)
            result, confidence = predict_diagnosis_with_confidence(inputs)
            diagnosis = diagnoses.get(result, 'Unknown')
            conclusion = f"Based on the provided inputs, the diagnosis is categorized as '{diagnosis}' with a confidence of {confidence:.2f}%. Please consult a specialist for further evaluation and management."
            st.markdown(f"<h2 style='text-align: center; color: {diagnosis_color};'>{diagnosis} ({confidence:.2f}%)</h2>", unsafe_allow_html=True)
            st.markdown(f"### Conclusion\n{conclusion}")

            # Show corresponding image
            if diagnosis == "Negative":
                st.image("Images/healthy.png", caption="Thyroid appears healthy", use_column_width=True)
            elif diagnosis == "Hypothyroid":
                st.image("Images/hypothyroid.png", caption="Hypothyroidism Condition", use_column_width=True)
            elif diagnosis == "Hyperthyroid":
                st.image("Images/hyperthyroid.png", caption="Hyperthyroidism Condition", use_column_width=True)

            df = pd.DataFrame([[name, email, diagnosis, datetime.now().strftime('%Y-%m-%d %H:%M:%S')]], columns=['Name', 'Email', 'Diagnosis', 'Date'])
            df.to_csv('diagnosis_history.csv', mode='a', header=not os.path.exists('diagnosis_history.csv'), index=False)

            pdf_file = generate_pdf(name, diagnosis, confidence, conclusion)
            with open(pdf_file, "rb") as f:
                st.download_button("Download PDF Report", f, file_name=pdf_file)

            if send_email_report(email, pdf_file):
                st.success("Report sent to email!")

            if diagnosis == "Hypothyroid":
                st.info("Recommended: Consult an endocrinologist. [Learn more](https://www.mayoclinic.org/diseases-conditions/hypothyroidism/)")
            elif diagnosis == "Hyperthyroid":
                st.info("Recommended: Schedule a specialist consultation. [Learn more](https://www.mayoclinic.org/diseases-conditions/hyperthyroidism/)")

    if os.path.exists('diagnosis_history.csv'):
        data = pd.read_csv('diagnosis_history.csv')
        st.subheader("📊 Diagnosis History Summary")
        st.bar_chart(data['Diagnosis'].value_counts())

if __name__ == '__main__':
    main()
